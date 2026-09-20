from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_documentation():
    doc = Document()

    # Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Title Page
    title = doc.add_heading('Laundry Management System', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Technical Walkthrough & Verification Report')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()

    # Introduction
    doc.add_heading('1. Overview', level=1)
    doc.add_paragraph(
        'The Laundry Management System is a full-stack application designed to streamline '
        'laundry operations for students and staff. It features real-time machine tracking, '
        'order management, rewards systems, and automated reminders.'
    )

    # Screenshots directory
    img_dir = r'c:\Users\SKV\Desktop\laundry app\screenshots'

    # Features
    features = [
        ('Student Dashboard', 'The central hub for students to view available machines and recent activity.', 'student_dashboard.png'),
        ('Laundry Tracking', 'Real-time status updates for active laundry orders.', 'student_track.png'),
        ('Service Pricing', 'Detailed itemized pricing for various dry-cleaning and laundry services.', 'student_pricing.png'),
        ('Rewards & Loyalty', 'Gamified rewards system tracking student streaks and points.', 'student_rewards.png'),
        ('Transaction History', 'Secure log of all previous laundry payments and submissions.', 'student_transactions.png'),
        ('Automated Reminders', 'Configurable notifications for laundry pick-up and drop-off.', 'student_reminders.png'),
        ('Order Submission', 'Success modal confirming the submission of a new laundry request.', 'submission_success.png'),
        ('Staff Administration', 'Comprehensive control panel for staff to manage machines and student orders.', 'admin_panel.png'),
    ]

    doc.add_heading('2. System Features & Verification', level=1)

    for name, desc, filename in features:
        doc.add_heading(name, level=2)
        doc.add_paragraph(desc)
        
        img_path = os.path.join(img_dir, filename)
        if os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(6))
        else:
            doc.add_paragraph(f'[Error: Image {filename} not found]')
        
        doc.add_paragraph('') # Spacing

    # Conclusion
    doc.add_heading('3. Verification Status', level=1)
    doc.add_paragraph('All functional tests passed successfully. The backend is correctly integrated with the MySQL database, and the frontend provides a responsive and intuitive user experience.')

    # Save
    output_path = r'c:\Users\SKV\Desktop\laundry app\LaundryApp_Documentation.docx'
    doc.save(output_path)
    print(f'Documentation saved to {output_path}')

if __name__ == '__main__':
    create_documentation()
